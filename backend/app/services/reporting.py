"""Reports and KPIs.

Formats: json, csv, md, xlsx (openpyxl), docx (python-docx), pdf (reportlab, optional).
Report kinds: ciso_brief, soa, risk_register, gap, roi (DORA Register of Information), incidents, vulnerabilities,
kpi, policies, assets, evidence_index. Each generated report is stored under NEXUS_REPORT_DIR with a hash.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
from datetime import UTC, datetime, timedelta
from pathlib import Path

from sqlalchemy.orm import Session

from app.db.models import (
    Asset,
    Control,
    Evidence,
    GapAssessment,
    Incident,
    Organization,
    Policy,
    Report,
    Risk,
    Task,
    Vendor,
    Vulnerability,
)
from app.services.gap_analysis import coverage_by_framework
from app.services.risk_engine import heatmap

REPORT_DIR = Path(os.getenv("NEXUS_REPORT_DIR", "./data/reports"))

ROI_COLUMNS = ["entity_name", "entity_lei", "provider_name", "provider_lei", "provider_country", "type_of_ict_service",
               "supports_critical_or_important_function", "tier", "contract_reference", "contract_start", "contract_end",
               "notice_period_days", "data_location", "personal_data", "sub_outsourcing", "subcontractors", "substitutability",
               "exit_strategy", "last_assessment", "residual_score"]


# ----------------------------------------------------------------------------- data builders
def kpis(db: Session, org_id: str) -> dict:
    now = datetime.now(UTC)
    risks = db.query(Risk).filter_by(org_id=org_id).all()
    open_risks = [r for r in risks if r.treatment not in ("closed",)]
    incidents = db.query(Incident).filter_by(org_id=org_id).all()
    open_inc = [i for i in incidents if i.status != "closed"]
    vendors = db.query(Vendor).filter_by(org_id=org_id).all()
    evidence = db.query(Evidence).filter_by(org_id=org_id).all()
    controls = db.query(Control).filter_by(org_id=org_id).all()
    vulns = db.query(Vulnerability).filter_by(org_id=org_id, status="open").all()
    tasks = db.query(Task).filter(Task.org_id == org_id, Task.status != "done").all()
    policies = db.query(Policy).filter_by(org_id=org_id).all()
    overdue_deadlines = []
    for i in open_inc:
        for d in i.notification_deadlines or []:
            try:
                if datetime.fromisoformat(d["due"]) < now:
                    overdue_deadlines.append({"incident": i.title, **d})
            except (ValueError, TypeError, KeyError):
                pass
    return {
        "risks": {"total": len(risks), "open": len(open_risks), "high": len([r for r in open_risks if r.level in ("high", "critical")]),
                  "critical_avg_score": round(sum((r.residual_score or r.inherent_score) for r in open_risks if r.level == "critical") / max(1, len([r for r in open_risks if r.level == "critical"])), 1),
                  "accepted": len([r for r in risks if r.treatment == "accept"]), "heatmap": heatmap(open_risks)},
        "incidents": {"open": len(open_inc), "dora_major": len([i for i in incidents if i.dora_major]),
                      "overdue_notifications": len(overdue_deadlines), "mttr_days": _mttr(incidents)},
        "vendors": {"total": len(vendors), "tier1": len([v for v in vendors if v.tier == 1]),
                    "critical_function": len([v for v in vendors if v.supports_critical_function]),
                    "no_exit_strategy": len([v for v in vendors if v.supports_critical_function and not v.exit_strategy]),
                    "assessment_overdue": len([v for v in vendors if v.next_assessment_at and v.next_assessment_at < now])},
        "evidence": {"total": len(evidence), "pass": len([e for e in evidence if e.result == "pass"]),
                     "fail": len([e for e in evidence if e.result == "fail"]), "unmapped": len([e for e in evidence if not e.control_id])},
        "controls": {"total": len(controls), "implemented": len([c for c in controls if c.status == "implemented"]),
                     "partial": len([c for c in controls if c.status == "partially_implemented"]),
                     "avg_maturity": round(sum(c.maturity for c in controls) / len(controls), 2) if controls else 0},
        "vulnerabilities": {"open": len(vulns), "critical": len([v for v in vulns if v.severity == "critical"]),
                            "sla_breached": len([v for v in vulns if v.due_at and v.due_at < now])},
        "tasks": {"open": len(tasks), "overdue": len([t for t in tasks if t.due_at and t.due_at < now]),
                  "by_status": {s: len([t for t in tasks if t.status == s]) for s in ("backlog", "todo", "doing", "review")}},
        "policies": {"total": len(policies), "approved": len([p for p in policies if p.status == "approved"]),
                     "review_due": len([p for p in policies if p.next_review_at and p.next_review_at < now + timedelta(days=30)])},
        "frameworks": coverage_by_framework(db, org_id),
        "overdue_deadlines": overdue_deadlines,
    }


def _mttr(incidents: list[Incident]) -> float | None:
    closed = [i for i in incidents if i.status in ("resolved", "closed")]
    if not closed:
        return None
    return round(sum((i.updated_at - i.detected_at).total_seconds() for i in closed) / len(closed) / 86400, 1)


def soa_rows(db: Session, org_id: str, framework_id: str = "iso27001-2022") -> list[dict]:
    rows = []
    for c in db.query(Control).filter_by(org_id=org_id).all():
        if c.requirement.framework_id != framework_id:
            continue
        ev_count = db.query(Evidence).filter_by(control_id=c.id).count()
        rows.append({"ref": c.requirement.ref, "title": c.requirement.title, "domain": c.requirement.domain, "applicable": c.applicable,
                     "justification": c.justification, "status": c.status, "maturity": c.maturity, "owner": c.owner or "",
                     "evidence_count": ev_count, "last_reviewed": c.last_reviewed_at.isoformat() if c.last_reviewed_at else ""})
    rows.sort(key=lambda r: [int(p) if p.isdigit() else p for p in str(r["ref"]).replace("A.", "").split(".")])
    return rows


def risk_rows(db: Session, org_id: str) -> list[dict]:
    return [{"id": r.id, "title": r.title, "category": r.category, "threat": r.threat, "vulnerability": r.vulnerability,
             "likelihood": r.likelihood, "impact": r.impact, "inherent_score": r.inherent_score, "residual_score": r.residual_score,
             "level": r.level, "treatment": r.treatment, "owner": r.owner or "", "review_at": r.review_at.isoformat() if r.review_at else "",
             "external": f"{r.external_system}:{r.external_id}" if r.external_system else ""}
            for r in db.query(Risk).filter_by(org_id=org_id).order_by(Risk.inherent_score.desc()).all()]


def roi_rows(db: Session, org_id: str) -> list[dict]:
    org = db.get(Organization, org_id)
    rows = []
    for v in db.query(Vendor).filter_by(org_id=org_id).order_by(Vendor.tier).all():
        rows.append({"entity_name": org.name if org else "", "entity_lei": (org.settings or {}).get("lei", "") if org else "",
                     "provider_name": v.name, "provider_lei": v.lei or "", "provider_country": v.country, "type_of_ict_service": v.service_type,
                     "supports_critical_or_important_function": "Y" if v.supports_critical_function else "N", "tier": v.tier,
                     "contract_reference": v.contract_ref or "", "contract_start": v.contract_start.date().isoformat() if v.contract_start else "",
                     "contract_end": v.contract_end.date().isoformat() if v.contract_end else "", "notice_period_days": v.notice_period_days or "",
                     "data_location": v.data_location, "personal_data": "Y" if v.personal_data else "N", "sub_outsourcing": "Y" if v.sub_outsourcing else "N",
                     "subcontractors": "; ".join(s if isinstance(s, str) else s.get("name", "") for s in (v.subcontractors or [])),
                     "substitutability": v.substitutable, "exit_strategy": "Y" if v.exit_strategy else "N",
                     "last_assessment": v.last_assessment_at.date().isoformat() if v.last_assessment_at else "", "residual_score": v.residual_score or ""})
    return rows


def incident_rows(db: Session, org_id: str) -> list[dict]:
    return [{"id": i.id, "title": i.title, "detected_at": i.detected_at.isoformat(), "status": i.status, "severity": i.severity, "source": i.source,
             "dora_major": i.dora_major, "nis2_significant": i.nis2_significant, "personal_data_breach": i.personal_data_breach,
             "deadlines": "; ".join(f"{d['regime']} {d['due']}" for d in (i.notification_deadlines or [])), "ticket": i.ticket_ref or ""}
            for i in db.query(Incident).filter_by(org_id=org_id).order_by(Incident.detected_at.desc()).all()]


def vuln_rows(db: Session, org_id: str) -> list[dict]:
    return [{"cve": v.cve or "", "title": v.title, "severity": v.severity, "cvss": v.cvss, "asset": v.asset_name, "status": v.status,
             "detected_at": v.detected_at.isoformat(), "due_at": v.due_at.isoformat() if v.due_at else "", "source": v.source}
            for v in db.query(Vulnerability).filter_by(org_id=org_id).order_by(Vulnerability.severity).all()]


def asset_rows(db: Session, org_id: str) -> list[dict]:
    return [{"name": a.name, "type": a.asset_type, "owner": a.owner or "", "classification": a.classification, "criticality": a.criticality,
             "critical_function": a.supports_critical_function, "location": a.location, "source": f"{a.external_system}:{a.external_id}" if a.external_system else ""}
            for a in db.query(Asset).filter_by(org_id=org_id).all()]


def policy_rows(db: Session, org_id: str) -> list[dict]:
    return [{"code": p.code, "title": p.title, "kind": p.doc_kind, "status": p.status, "version": p.current_version, "owner": p.owner or "",
             "next_review": p.next_review_at.date().isoformat() if p.next_review_at else "",
             "requirements": "; ".join(f"{r['framework_id']} {r['ref']}" for r in (p.requirement_refs or []))}
            for p in db.query(Policy).filter_by(org_id=org_id).order_by(Policy.code).all()]


def evidence_rows(db: Session, org_id: str) -> list[dict]:
    return [{"title": e.title, "source": e.source, "source_ref": e.source_ref, "result": e.result, "collected_at": e.collected_at.isoformat(),
             "control": e.control.requirement.ref if e.control else "", "sha256": e.sha256, "file": e.file_path or ""}
            for e in db.query(Evidence).filter_by(org_id=org_id).order_by(Evidence.collected_at.desc()).all()]


def gap_rows(db: Session, org_id: str, assessment_id: str | None = None) -> list[dict]:
    q = db.query(GapAssessment).filter_by(org_id=org_id)
    ga = q.filter_by(id=assessment_id).first() if assessment_id else q.order_by(GapAssessment.created_at.desc()).first()
    if not ga:
        return []
    from app.db.models import Requirement

    rows = []
    for it in ga.items:
        req = db.get(Requirement, it.requirement_id)
        rows.append({"assessment": ga.title, "coverage_pct": ga.coverage_pct, "ref": req.ref if req else "", "title": req.title if req else "",
                     "domain": req.domain if req else "", "status": it.current_status, "severity": it.gap_severity, "recommendation": it.recommendation,
                     "task_id": it.task_id or ""})
    return rows


def ciso_brief_markdown(db: Session, org_id: str) -> str:
    org = db.get(Organization, org_id)
    k = kpis(db, org_id)
    lines = [f"# CISO Brief - {org.name if org else org_id}", f"_Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC_", "",
             "## Posture at a glance",
             f"- Risks: {k['risks']['open']} open, {k['risks']['high']} high or critical, {k['risks']['accepted']} accepted",
             f"- Incidents: {k['incidents']['open']} open, {k['incidents']['dora_major']} DORA-major to date, {k['incidents']['overdue_notifications']} overdue regulatory notifications",
             f"- Third parties: {k['vendors']['total']} registered, {k['vendors']['critical_function']} supporting critical functions, {k['vendors']['no_exit_strategy']} of those without exit strategy",
             f"- Vulnerabilities: {k['vulnerabilities']['open']} open, {k['vulnerabilities']['critical']} critical, {k['vulnerabilities']['sla_breached']} past SLA",
             f"- Controls: {k['controls']['implemented']}/{k['controls']['total']} implemented, average maturity {k['controls']['avg_maturity']}/5",
             f"- Evidence: {k['evidence']['total']} artefacts, {k['evidence']['fail']} failing, {k['evidence']['unmapped']} unmapped",
             f"- Work: {k['tasks']['open']} open tasks, {k['tasks']['overdue']} overdue; policies approved {k['policies']['approved']}/{k['policies']['total']}", "",
             "## Framework coverage"]
    for f in k["frameworks"]:
        lines.append(f"- {f['name']}: {f['coverage_pct']}%")
    if k["overdue_deadlines"]:
        lines += ["", "## Regulatory deadlines passed"]
        lines += [f"- {d['regime']} for '{d['incident']}' (due {d['due']})" for d in k["overdue_deadlines"]]
    top = [r for r in db.query(Risk).filter_by(org_id=org_id).order_by(Risk.inherent_score.desc()).limit(5).all()]
    if top:
        lines += ["", "## Top risks"] + [f"- [{r.level}] {r.title} (score {r.residual_score or r.inherent_score}, {r.treatment})" for r in top]
    lines += ["", "## Decision requested this week",
              "- " + (f"Approve treatment for the {k['risks']['high']} high/critical risks and the exit strategies for {k['vendors']['no_exit_strategy']} critical providers"
                      if k['risks']['high'] or k['vendors']['no_exit_strategy'] else "No escalation required; continue the ISMS plan")]
    return "\n".join(lines)


# ----------------------------------------------------------------------------- renderers
def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict], sheet: str = "Report") -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        for r in rows:
            ws.append([json.dumps(r[h], default=str) if isinstance(r[h], (list, dict)) else r[h] for h in headers])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_docx(title: str, markdown: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(title, 0)
    for line in markdown.splitlines():
        if line.startswith("# "):
            d.add_heading(line[2:], 1)
        elif line.startswith("## "):
            d.add_heading(line[3:], 2)
        elif line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            d.add_paragraph(line)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def to_pdf(title: str, markdown: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for line in markdown.splitlines():
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Heading1"]))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.strip():
            story.append(Paragraph(line.replace("- ", "• ", 1) if line.startswith("- ") else line, styles["BodyText"]))
    doc.build(story)
    return out.getvalue()


BUILDERS = {"soa": soa_rows, "risk_register": risk_rows, "roi": roi_rows, "incidents": incident_rows, "vulnerabilities": vuln_rows,
            "assets": asset_rows, "policies": policy_rows, "evidence_index": evidence_rows, "gap": gap_rows}


def generate(db: Session, org_id: str, kind: str, fmt: str, actor: str, **kw) -> tuple[bytes, str, Report]:
    """Returns (content bytes, media type, Report row)."""
    title = f"{kind.replace('_', ' ').title()} - {datetime.now(UTC):%Y-%m-%d}"
    if kind in ("ciso_brief", "kpi"):
        md = ciso_brief_markdown(db, org_id) if kind == "ciso_brief" else "```json\n" + json.dumps(kpis(db, org_id), indent=2, default=str) + "\n```"
        if fmt == "md":
            data, media = md.encode(), "text/markdown"
        elif fmt == "json":
            data, media = json.dumps({"markdown": md, "kpis": kpis(db, org_id)}, default=str).encode(), "application/json"
        elif fmt == "docx":
            data, media = to_docx(title, md), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "pdf":
            data, media = to_pdf(title, md), "application/pdf"
        else:
            raise ValueError("format not supported for this report")
    else:
        rows = BUILDERS[kind](db, org_id, **kw) if kind == "gap" else BUILDERS[kind](db, org_id)
        if fmt == "csv":
            data, media = to_csv(rows), "text/csv"
        elif fmt == "xlsx":
            data, media = to_xlsx(rows, kind), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif fmt == "json":
            data, media = json.dumps(rows, default=str).encode(), "application/json"
        elif fmt == "md":
            md = f"# {title}\n\n" + ("| " + " | ".join(rows[0].keys()) + " |\n|" + "---|" * len(rows[0]) + "\n" +
                                   "\n".join("| " + " | ".join(str(v) for v in r.values()) + " |" for r in rows) if rows else "_no rows_")
            data, media = md.encode(), "text/markdown"
        else:
            raise ValueError("format not supported for this report")
    REPORT_DIR.joinpath(org_id).mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(data).hexdigest()
    path = REPORT_DIR / org_id / f"{kind}-{datetime.now(UTC):%Y%m%d-%H%M%S}.{fmt}"
    path.write_bytes(data)
    rep = Report(org_id=org_id, kind=kind, fmt=fmt, title=title, file_path=str(path), generated_by=actor, sha256=digest,
                 content=data.decode("utf-8", errors="ignore") if fmt in ("md", "json", "csv") else "")
    db.add(rep)
    db.commit()
    return data, media, rep
