"""Document library: imports a local folder tree (e.g. a synced Google Drive folder) into the repository.

The folder is mounted read-only at NEXUS_LIBRARY_DIR. Originals stay where they are, with their folder
structure and names; each file becomes a global KnowledgeDocument (doc_type "library") holding the
extracted text, its relative path and one tag per parent folder. The library is local-only: it is
never committed or baked into images, because it can hold licensed third-party material.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import warnings
import zipfile
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.logging import get_logger
from app.db.models import KnowledgeDocument

log = get_logger(__name__)
logging.getLogger("pypdf").setLevel(logging.ERROR)  # malformed-PDF chatter would flood the import log
warnings.filterwarnings("ignore", module="openpyxl")
DOC_TYPE = "library"
MAX_TEXT = 1_000_000  # characters kept per document
MAX_PDF_PAGES = 600
TEXT_TYPES = {".pdf", ".docx", ".dotx", ".xlsx", ".xlsm", ".pptx", ".csv", ".txt", ".md", ".xmind"}


def library_dir() -> Path | None:
    d = os.getenv("NEXUS_LIBRARY_DIR")
    return Path(d) if d and Path(d).is_dir() else None


def resolve(rel: str) -> Path | None:
    """Absolute path of a library file, refusing anything outside the library root."""
    root = library_dir()
    if not root:
        return None
    p = (root / rel).resolve()
    return p if p.is_file() and root.resolve() in p.parents else None


def _pdf(p: Path) -> str:
    from pypdf import PdfReader

    r = PdfReader(str(p))
    return "\n".join((pg.extract_text() or "") for pg in r.pages[:MAX_PDF_PAGES])


def _docx(p: Path) -> str:
    import docx

    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _xlsx(p: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _pptx(p: Path) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        out.append(f"## Slide {i}")
        out += [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    return "\n".join(out)


def _xmind(p: Path) -> str:
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        if "content.json" in names:
            out: list[str] = []

            def walk(t: dict, depth: int) -> None:
                if t.get("title"):
                    out.append("  " * depth + "- " + t["title"])
                for kids in (t.get("children") or {}).values():
                    for k in kids:
                        walk(k, depth + 1)

            for sheet in json.loads(z.read("content.json")):
                out.append(f"## {sheet.get('title', '')}")
                walk(sheet.get("rootTopic", {}), 0)
            return "\n".join(out)
        if "content.xml" in names:
            xml = z.read("content.xml").decode("utf-8", "ignore")
            return "\n".join(re.findall(r"<title[^>]*>([^<]+)</title>", xml))
    return ""


def _csv(p: Path) -> str:
    raw = p.read_bytes().decode("utf-8", "ignore")
    return "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(raw)))


EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".dotx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
              ".xmind": _xmind, ".csv": _csv, ".txt": lambda p: p.read_text("utf-8", "ignore"),
              ".md": lambda p: p.read_text("utf-8", "ignore")}


def _language(text: str) -> str:
    sample = f" {text[:20000].lower()} "
    pt = sum(sample.count(w) for w in (" de ", " que ", " não ", " para ", " da ", " dos "))
    en = sum(sample.count(w) for w in (" the ", " and ", " of ", " to ", " is ", " for "))
    return "pt-PT" if pt > en else "en"


def import_library(db: Session, limit: int | None = None) -> dict:
    """Upsert one document per file. Unchanged files (same size, mtime and path) are skipped."""
    root = library_dir()
    if not root:
        return {"error": "NEXUS_LIBRARY_DIR is not set or not a directory"}
    existing = {d.file_path: d for d in db.query(KnowledgeDocument).filter_by(doc_type=DOC_TYPE, org_id=None)}
    stats = {"files": 0, "imported": 0, "unchanged": 0, "no_text": 0, "failed": 0, "removed": 0}
    seen: set[str] = set()
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name.startswith(".") or path.name == "Icon\r":
            continue
        rel = path.relative_to(root).as_posix()
        seen.add(rel)
        stats["files"] += 1
        st = path.stat()
        # Change marker (not a content hash): cheap to compare on re-import of thousands of files.
        marker = hashlib.sha256(f"{rel}:{st.st_size}:{int(st.st_mtime)}".encode()).hexdigest()
        doc = existing.get(rel)
        if doc and doc.sha256 == marker:
            stats["unchanged"] += 1
            continue
        ext = path.suffix.lower()
        text = ""
        if ext in EXTRACTORS:
            try:
                text = re.sub(r"[ \t]+\n", "\n", EXTRACTORS[ext](path)).replace("\x00", "").strip()[:MAX_TEXT]
            except Exception as exc:  # malformed or protected files are listed without text
                stats["failed"] += 1
                log.warning("library.extract_failed", path=rel, error=str(exc)[:200])
        folders = list(Path(rel).parent.parts)
        tags = folders + ([] if text else ["sem-texto"])
        doc = doc or KnowledgeDocument(org_id=None, doc_type=DOC_TYPE, file_path=rel)
        doc.title, doc.text, doc.sha256, doc.tags = path.name[:512], text, marker, tags
        doc.language = _language(text) if text else "pt-PT"
        doc.source_url = ""
        db.add(doc)
        stats["imported"] += 1
        if not text:
            stats["no_text"] += 1
        if stats["imported"] % 50 == 0:
            db.commit()
            log.info("library.progress", **stats)
        if limit and stats["imported"] >= limit:
            break
    if not limit:
        for rel, doc in existing.items():
            if rel not in seen:
                db.delete(doc)
                stats["removed"] += 1
    db.commit()
    log.info("library.import", **stats)
    return stats


def library_tree(db: Session) -> list[dict]:
    """Folder paths with document counts (each folder also counts documents in its subfolders)."""
    counts: dict[str, int] = {}
    for (fp,) in db.query(KnowledgeDocument.file_path).filter_by(doc_type=DOC_TYPE, org_id=None):
        parts = Path(fp or "").parent.parts
        for i in range(1, len(parts) + 1):
            key = "/".join(parts[:i])
            counts[key] = counts.get(key, 0) + 1
    return [{"path": k, "name": k.rsplit("/", 1)[-1], "depth": k.count("/"), "documents": v} for k, v in sorted(counts.items())]
