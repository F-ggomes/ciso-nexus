"""Evidence ingestion: connector artefacts and uploaded files (PDF, DOCX, XLSX, CSV, logs, images, JSON).

Files are stored under NEXUS_EVIDENCE_DIR/<org_id>/<sha256>.<ext>, de-duplicated by hash, with text extracted
for search and AI mapping. Uploads are limited by size and allowed MIME types (see ALLOWED), and the stored
name never comes from the client.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from app.connectors.base import EvidenceArtefact
from app.db.models import Control, Evidence, Requirement
from app.services.refs import parent_refs

EVIDENCE_DIR = Path(os.getenv("NEXUS_EVIDENCE_DIR", "./data/evidence"))
MAX_UPLOAD_BYTES = int(os.getenv("NEXUS_MAX_UPLOAD_MB", "50")) * 1024 * 1024
ALLOWED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/csv": "csv", "text/plain": "txt", "application/json": "json", "text/x-log": "log",
    "image/png": "png", "image/jpeg": "jpg",
}


def extract_text(data: bytes, ext: str) -> str:
    try:
        if ext == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages[:200])
        if ext == "docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs)
        if ext == "xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:10]:
                for row in ws.iter_rows(max_row=2000, values_only=True):
                    lines.append("\t".join("" if c is None else str(c) for c in row))
            return "\n".join(lines)
        if ext in ("csv", "txt", "log", "json"):
            return data.decode("utf-8", errors="replace")[:2_000_000]
    except Exception as exc:  # extraction is best effort
        return f"[extraction failed: {exc}]"
    return ""


def store_upload(db: Session, org_id: str, filename: str, mime: str, data: bytes, title: str | None, control_id: str | None,
                 result: str = "info") -> Evidence:
    if mime not in ALLOWED:
        raise ValueError(f"Unsupported file type {mime}")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File too large")
    ext = ALLOWED[mime]
    digest = hashlib.sha256(data).hexdigest()
    folder = EVIDENCE_DIR / org_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{digest}.{ext}"
    if not path.exists():
        path.write_bytes(data)
    ev = Evidence(org_id=org_id, control_id=control_id, source="upload", source_ref=filename[:255], title=(title or filename)[:512],
                  payload={"original_name": filename, "ext": ext}, file_path=str(path), mime_type=mime, size_bytes=len(data),
                  extracted_text=extract_text(data, ext)[:500_000], sha256=digest, result=result)
    db.add(ev)
    db.commit()
    return ev


def store_artefacts(db: Session, org_id: str, artefacts: list[EvidenceArtefact], auto_link: bool = True) -> list[Evidence]:
    """Persist connector artefacts; skip unchanged payloads (same source_ref + sha256); auto-link by candidate refs."""
    out: list[Evidence] = []
    for a in artefacts:
        dup = db.query(Evidence).filter_by(org_id=org_id, source=a.source, source_ref=a.source_ref, sha256=a.sha256).first()
        if dup:
            dup.collected_at = datetime.now(timezone.utc)
            continue
        control_id = None
        if auto_link and a.candidate_control_refs:
            control_id = _match_control(db, org_id, a.candidate_control_refs)
        ev = Evidence(org_id=org_id, control_id=control_id, source=a.source, source_ref=a.source_ref, title=a.title,
                      payload=json.loads(json.dumps(a.payload, default=str)), sha256=a.sha256, result=a.result,
                      collected_at=a.collected_at, extracted_text=json.dumps(a.payload, default=str)[:200_000])
        db.add(ev)
        out.append(ev)
    db.commit()
    return out


def _match_control(db: Session, org_id: str, refs: list[str]) -> str | None:
    for ref in refs:
        fw_hint = None
        if ref.startswith("CIS-"):
            fw_hint, ref = "cis-v8.1", ref[4:]
        elif ref.startswith("DORA-"):
            fw_hint, ref = "dora", ref[5:].replace("Art", "Art. ")
        for candidate in parent_refs(ref):
            q = db.query(Requirement).filter(Requirement.ref == candidate)
            if fw_hint:
                q = q.filter(Requirement.framework_id == fw_hint)
            for req in q.all():
                ctrl = db.query(Control).filter_by(org_id=org_id, requirement_id=req.id).first()
                if ctrl:
                    return ctrl.id
    return None


def candidate_controls(db: Session, org_id: str, limit: int = 40) -> list[dict]:
    rows = db.query(Control).filter_by(org_id=org_id).limit(limit * 10).all()
    return [{"control_id": c.id, "framework": c.requirement.framework_id, "control_ref": c.requirement.ref,
             "title": c.requirement.title} for c in rows][:limit]
